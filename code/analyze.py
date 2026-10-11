"""정류장에서 다시 재는 저상버스 — 서울 시내버스.

저상버스 정책 성과는 **보유율(%)** 로 보고된다. 그런데 휠체어를 탄 사람은 차량을 고르지
못한다. 정류장에 서서 오는 버스를 받을 뿐이다. 그래서 같은 보유율이라도 저상차가 특정
노선에 몰려 있으면 정류장마다 기다리는 시간이 갈린다.

이 스크립트는 정류장마다 두 가지를 잰다.

    전체 유효 배차 h_all = 1 / Σ_r (1 / h_r)
    저상 유효 배차 h_lf  = 1 / Σ_r (p_r / h_r)

h_r 은 노선 r 의 배차간격(분), p_r 은 그 노선의 저상버스 보유율이다. 저상차는 노선의
차량 운용에 섞여 돌므로, 한 번의 도착이 저상일 확률을 p_r 로 둔다.

**배수 = h_lf / h_all** 이 1 이면 저상이 아닌 버스와 똑같이 기다린다는 뜻이고,
3 이면 세 배를 더 기다린다는 뜻이다.

자료
  d01  서울시 저상버스 도입 노선 및 노선별 보유율   서울특별시 · 열린데이터광장 · 2025-04-25
  d02  서울시 버스노선별 정류소 정보                서울특별시 · 열린데이터광장 · 2026-09-02
  d03  서울시 버스 인가현황                          서울특별시 · TOPIS · 2026-09

실행: python projects/koreashe_traffic_2026/analyze.py
"""
from __future__ import annotations

import csv
import re
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parents[1]
data = root / "data"
out_dir = root / "out"

# 집계에 든 노선의 운행 빈도가 정류장 전체의 몇 할을 덮는지. 이보다 낮으면
# 절대 분(分)을 믿을 수 없어 요약 통계에서 뺀다.
MIN_COVER = 0.7


def sheet(path: Path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return wb[wb.sheetnames[0]]


def norm(x) -> str:
    """노선번호 표기를 맞춘다.

    파일마다 다르게 적혀 있다. 숫자로 저장된 칸, A/B 로 갈린 노선(110A·110B 를
    보유율 표는 110 하나로 적는다), 맞춤버스의 '출근'·'퇴근' 꼬리표가 그렇다.
    """
    if x is None:
        return ""
    s = str(x).strip().replace(" ", "")
    if s.endswith(".0"):
        s = s[:-2]
    s = re.sub(r"(출근|퇴근)$", "", s)
    s = re.sub(r"(?<=\d)[AB]$", "", s)
    return s


def load_lowfloor() -> dict[str, dict]:
    """노선번호 -> 인가/저상 대수. A·B 로 갈린 노선은 합쳐 둔다."""
    agg: dict[str, dict] = {}
    for r in sheet(data / "d01_저상버스_보유율.xlsx").iter_rows(min_row=2, values_only=True):
        if r[0] is None:
            continue
        k = norm(r[0])
        a = agg.setdefault(k, {"인가": 0, "저상": 0})
        a["인가"] += r[1] or 0
        a["저상"] += r[2] or 0
    for v in agg.values():
        v["보유율"] = v["저상"] / v["인가"] if v["인가"] else 0.0
    return agg


def load_routes() -> dict[str, dict]:
    """노선번호 -> 유형·배차·첫차·막차. 공동배차 노선은 대수를 더한다."""
    agg: dict[str, dict] = {}
    for r in sheet(data / "d03_인가현황_202609.xlsx").iter_rows(min_row=2, values_only=True):
        if r[1] is None:
            continue
        try:
            lo, hi = float(r[11]), float(r[12])
        except (TypeError, ValueError):
            continue
        if lo <= 0 or hi <= 0:
            continue
        k = norm(r[1])
        a = agg.setdefault(k, {"유형": r[2], "인가": 0, "최소배차": lo,
                               "최대배차": hi, "첫차": r[13], "막차": r[14]})
        a["인가"] += r[5] or 0
        a["최소배차"] = min(a["최소배차"], lo)
        a["최대배차"] = max(a["최대배차"], hi)
    return agg


# 서울시 정류소 고유번호(NODE_ID)는 앞 세 자리가 자치구를 가리킨다. 100 종로구부터
# 124 강동구까지 스물다섯이고, 2 로 시작하면 경기·인천이다. 서울 시내버스가 경기 땅을
# 지나기도 해서, 서울 안 정류장만 남기려면 이 자리를 본다.
GU = {100: "종로구", 101: "중구", 102: "용산구", 103: "성동구", 104: "광진구",
      105: "동대문구", 106: "중랑구", 107: "성북구", 108: "강북구", 109: "도봉구",
      110: "노원구", 111: "은평구", 112: "서대문구", 113: "마포구", 114: "양천구",
      115: "강서구", 116: "구로구", 117: "금천구", 118: "영등포구", 119: "동작구",
      120: "관악구", 121: "서초구", 122: "강남구", 123: "송파구", 124: "강동구"}


def gu_of(node: str) -> str | None:
    try:
        return GU.get(int(str(node)[:3]))
    except ValueError:
        return None


def load_stops():
    """NODE_ID -> (이름, ARS, 경도, 위도) 와 NODE_ID -> {노선번호}. 서울 안만."""
    info, served = {}, defaultdict(set)
    for r in sheet(data / "d02_노선별정류소.xlsx").iter_rows(min_row=2, values_only=True):
        if r[3] is None or r[6] is None:
            continue
        node = str(r[3]).strip()
        if gu_of(node) is None:
            continue
        info.setdefault(node, (r[5], r[4], float(r[6]), float(r[7])))
        served[node].add(norm(r[1]))
    return info, served


def effective(routes_at: set[str], routes: dict, low: dict, key: str):
    """정류장 하나의 유효 배차간격. key 는 '최소배차'(혼잡)·'최대배차'(한산)."""
    all_rate = lf_rate = cov_hit = cov_all = 0.0
    used = 0
    for rn in routes_at:
        r = routes.get(rn)
        if not r:
            continue
        h = r[key]
        cov_all += 1.0 / h
        if rn not in low:
            continue                       # 저상 자료가 없는 노선은 양쪽 모두에서 뺀다
        used += 1
        cov_hit += 1.0 / h
        all_rate += 1.0 / h
        lf_rate += low[rn]["보유율"] / h
    if used == 0 or all_rate == 0:
        return None
    return {
        "노선수": used,
        "커버리지": cov_hit / cov_all if cov_all else 0.0,
        "전체": 1 / all_rate,
        "저상": (1 / lf_rate) if lf_rate > 0 else float("inf"),
    }


def pct(xs, p):
    xs = sorted(xs)
    return xs[min(len(xs) - 1, int(len(xs) * p))]


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    out_dir.mkdir(exist_ok=True)

    low, routes = load_lowfloor(), load_routes()
    info, served = load_stops()
    both = set(low) & set(routes)
    print(f"저상 자료 {len(low)}개 노선 · 인가현황 {len(routes)}개 노선 · 맞물린 노선 {len(both)}개")
    print(f"서울 시내버스 전체 보유율 "
          f"{sum(v['저상'] for v in low.values()) / sum(v['인가'] for v in low.values()):.1%} "
          f"(저상 {sum(v['저상'] for v in low.values()):,} / 인가 {sum(v['인가'] for v in low.values()):,}대)")
    rates = sorted(low[k]["보유율"] for k in both)
    print(f"노선별 보유율 — 0% {sum(1 for x in rates if x == 0)}개 · "
          f"50% 미만 {sum(1 for x in rates if x < .5)}개 · "
          f"100% {sum(1 for x in rates if x >= .999)}개 · 중앙 {pct(rates, .5):.0%}")

    rows = []
    for node, rset in served.items():
        peak = effective(rset, routes, low, "최소배차")
        off = effective(rset, routes, low, "최대배차")
        if not peak or not off:
            continue
        name, ars, x, y = info[node]
        rows.append({
            "node": node, "ars": ars, "gu": gu_of(node),
            "name": name, "lon": x, "lat": y,
            "노선수": peak["노선수"], "커버리지": round(peak["커버리지"], 3),
            "혼잡_전체": round(peak["전체"], 2), "혼잡_저상": round(peak["저상"], 2),
            "한산_전체": round(off["전체"], 2), "한산_저상": round(off["저상"], 2),
            "배수_혼잡": round(peak["저상"] / peak["전체"], 3),
            "배수_한산": round(off["저상"] / off["전체"], 3),
        })

    good = [r for r in rows if r["커버리지"] >= MIN_COVER]
    dead = [r for r in good if r["혼잡_저상"] == float("inf")]
    live = [r for r in good if r["혼잡_저상"] != float("inf")]
    print(f"\n정류소 {len(info):,}곳 중 계산 {len(rows):,}곳 · "
          f"커버리지 {MIN_COVER:.0%} 이상 {len(good):,}곳")
    print(f"그중 저상버스가 한 대도 서지 않는 정류소 **{len(dead):,}곳** "
          f"({len(dead) / len(good):.1%})")

    r_peak = [r["배수_혼잡"] for r in live]
    print(f"\n배수(저상 ÷ 전체) — 중앙 {pct(r_peak, .5):.2f}배 · "
          f"상위 25% {pct(r_peak, .75):.2f}배 · 상위 10% {pct(r_peak, .9):.2f}배 · "
          f"상위 5% {pct(r_peak, .95):.2f}배 · 상위 1% {pct(r_peak, .99):.2f}배")

    peak_lf = [r["혼잡_저상"] for r in live]
    peak_all = [r["혼잡_전체"] for r in live]
    print(f"혼잡 시간대 유효 배차 — 전체 중앙 {pct(peak_all, .5):.1f}분 → "
          f"저상 중앙 {pct(peak_lf, .5):.1f}분 "
          f"· 저상 상위 10% {pct(peak_lf, .9):.1f}분")
    off_lf = [r["한산_저상"] for r in live]
    print(f"한산 시간대 유효 배차 — 저상 중앙 {pct(off_lf, .5):.1f}분 "
          f"· 상위 10% {pct(off_lf, .9):.1f}분")

    over = [r for r in live if r["혼잡_저상"] >= 30]
    print(f"\n혼잡 시간대에도 저상버스를 **30분 넘게** 기다리는 정류소 {len(over):,}곳 "
          f"({len(over) / len(good):.1%}) · 여기에 더해 아예 없는 곳 {len(dead):,}곳")

    with (out_dir / "stop_lowfloor.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        w.writeheader()
        for r in rows:
            w.writerow({k: ("" if v == float("inf") else v) for k, v in r.items()})
    print(f"\n-> {out_dir / 'stop_lowfloor.csv'}  ({len(rows):,}행)")


if __name__ == "__main__":
    main()
